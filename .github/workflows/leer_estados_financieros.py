"""
leer_estados_financieros.py
Matus Consulting · Tablero Grupo Argentilia · Pestaña Rentabilidad

Lee los estados de resultados mensuales de Contraloría directamente de la
carpeta de Google Drive y genera `datos_rentabilidad.json` para el tablero.

Reglas de lectura (validadas contra los 41 archivos ene–ago 2026):
  1. Solo se leen los archivos por unidad ("<UNIDAD> <MES> <AÑO>.xlsx").
     Los consolidados ("<MES> <AÑO>.xlsx") se ignoran: tienen tres formatos
     distintos en el año. El total del grupo se calcula sumando unidades.
  2. De cada unidad se toma SOLO el archivo más reciente: ya contiene todos
     los meses del año, el presupuesto y el comparativo del año anterior.
  3. Pestaña: "ESTADO DE RESULTADOS". Columnas por ENCABEZADO (fila 6),
     nunca por letra. Rubros por NOMBRE (columna A, con sinónimos), nunca
     por número de fila.
  4. El cheque promedio se CALCULA (ventas totales ÷ comensales). El renglón
     "PROMEDIO DE VENTA" del archivo acumulado trae #REF! en meses previos.

Uso:
  Drive (GitHub Actions):  python leer_estados_financieros.py
      requiere la variable de entorno GOOGLE_SA_JSON (credencial de la
      cuenta de servicio) y, opcional, CARPETA_EF_ID.
  Local (pruebas):         python leer_estados_financieros.py --local <carpeta>
"""

import io
import json
import os
import re
import sys
import unicodedata
from datetime import datetime, timezone

import openpyxl

CARPETA_EF_ID = os.environ.get("CARPETA_EF_ID", "1jDKGW4dhuy9JKvZ2g8Zuwx17qeE7kBcT")
SALIDA = os.environ.get("SALIDA_RENTABILIDAD", "datos_rentabilidad.json")

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
         "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

# Prefijo del archivo -> nombre en el tablero y fecha de apertura
UNIDADES = {
    "ARGENTILIA LEON": {"nombre": "Argentilia León", "apertura": None},
    "ARGENTILIA QRO": {"nombre": "Argentilia Querétaro", "apertura": None},
    "FRASCATI": {"nombre": "Frascati", "apertura": None},
    "MIKOH": {"nombre": "Mikoh", "apertura": "2025-09-25"},
    "MANOLO": {"nombre": "Manolo", "apertura": "2026-07-30"},
}

HOJA = "ESTADO DE RESULTADOS"
FILA_ENCABEZADOS = 6
FILA_MAX_CARATULA = 110   # los rubros de la carátula viven antes de esta fila

# Rubro del tablero -> nombres aceptados en la columna A (normalizados)
RUBROS = {
    "ventas_alimentos": ["INGRESOS ALIMENTOS"],
    "ventas_bebidas": ["INGRESOS BEBIDAS"],
    "otros_ingresos": ["OTROS INGRESOS"],
    "ventas": ["TOTAL DE INGRESOS"],
    "costo_venta": ["COSTO TOTAL"],
    "nomina": ["TOTAL NOMINA Y RELATIVOS", "TOTAL NOM Y RELAT", "TOTAL NOM. Y RELAT."],
    "gastos_departamentales": ["TOTAL GASTOS DEPARTAMENTALES"],
    "renta": ["TOTAL DE RENTA", "TOTAL RENTA", "TOTAL GASTOS DE OPERACION"],
    "utilidad_bruta_operacion": ["UTILIDAD BRUTA DE OPERACION"],
    "gastos_varios": ["TOTAL GASTOS VARIOS"],
    "utilidad": ["UTILIDAD DESPUES DE OTROS GASTOS", "UTILIDAD DESPUES DE IMPUESTOS"],
    "comensales": ["COMENSALES", "# COMENSALES"],
}
OPCIONALES = {"otros_ingresos"}   # Mikoh no lo trae en algunos meses

PATRON_ARCHIVO = re.compile(
    r"^(?P<unidad>.+?)\s+(?P<mes>" + "|".join(MESES) + r")\s+(?P<anio>\d{4})\.xlsx$",
    re.IGNORECASE)


def norm(texto):
    """Mayúsculas, sin acentos y sin espacios dobles."""
    if texto is None:
        return ""
    t = unicodedata.normalize("NFKD", str(texto))
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", t).strip().upper()


def numero(valor):
    """Número o None (celdas vacías, #REF!, #DIV/0!, texto)."""
    if isinstance(valor, (int, float)):
        return float(valor)
    return None


# ---------------------------------------------------------------- archivos

def clasificar(nombre):
    """(clave_unidad, índice_mes, año) o None si no es archivo por unidad."""
    m = PATRON_ARCHIVO.match(nombre.strip())
    if not m:
        return None
    unidad = norm(m.group("unidad"))
    if unidad not in UNIDADES:
        return None
    return unidad, MESES.index(norm(m.group("mes"))), int(m.group("anio"))


def elegir_mas_recientes(nombres):
    """De una lista de nombres de archivo, el más reciente por unidad."""
    elegidos, ignorados = {}, []
    for n in nombres:
        c = clasificar(n)
        if not c:
            ignorados.append(n)
            continue
        unidad, mes, anio = c
        if unidad not in elegidos or (anio, mes) > elegidos[unidad][1]:
            elegidos[unidad] = (n, (anio, mes))
    return {u: v[0] for u, v in elegidos.items()}, ignorados


def listar_drive(servicio, carpeta_id):
    archivos, token = [], None
    while True:
        r = servicio.files().list(
            q=f"'{carpeta_id}' in parents and trashed = false",
            fields="nextPageToken, files(id, name, modifiedTime)",
            pageSize=200, pageToken=token,
            supportsAllDrives=True, includeItemsFromAllDrives=True).execute()
        archivos += r.get("files", [])
        token = r.get("nextPageToken")
        if not token:
            return archivos


def descargar_drive(servicio, archivo_id):
    from googleapiclient.http import MediaIoBaseDownload
    buffer = io.BytesIO()
    descarga = MediaIoBaseDownload(buffer, servicio.files().get_media(fileId=archivo_id))
    listo = False
    while not listo:
        _, listo = descarga.next_chunk()
    buffer.seek(0)
    return buffer


def servicio_drive():
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    credencial = os.environ.get("GOOGLE_SA_JSON")
    if not credencial:
        sys.exit("Falta la variable GOOGLE_SA_JSON (credencial de la cuenta de servicio).")
    info = json.loads(credencial)
    cred = service_account.Credentials.from_service_account_info(
        info, scopes=["https://www.googleapis.com/auth/drive.readonly"])
    return build("drive", "v3", credentials=cred, cache_discovery=False)


# ---------------------------------------------------------------- lectura

def columnas_por_mes(encabezados, anio):
    """{mes: {'real': col, 'presupuesto': col, 'anterior': col}} (1-based)."""
    pos_total = {}
    for i, h in enumerate(encabezados):
        for m in MESES:
            if h == f"TOTAL {m} {anio}":
                pos_total.setdefault(m, i)
    inicios = sorted(pos_total.values())
    resultado = {}
    for m, i in pos_total.items():
        fin = next((x for x in inicios if x > i), len(encabezados))
        bloque = range(i + 1, fin)
        ppto = next((j for j in bloque if encabezados[j] == "PRESUPUESTO"), None)
        # comparativo año anterior; tolera "TOTAL NOVIEMBRE 2026" mal rotulado
        patron = re.compile(rf"^(TOTAL )?{m} \d{{4}}$")
        anterior = next((j for j in bloque if patron.match(encabezados[j])), None)
        resultado[m] = {"real": i + 1,
                        "presupuesto": ppto + 1 if ppto is not None else None,
                        "anterior": anterior + 1 if anterior is not None else None}
    return resultado


def columnas_acumulado(encabezados, anio):
    idx = {h: i + 1 for i, h in enumerate(encabezados) if h}
    return {"real": idx.get(f"ACUMULADO {anio}"), "anterior": idx.get(f"ACUMULADO {anio - 1}")}


def filas_rubros(ws, alertas, unidad):
    etiquetas = {}
    for r in range(1, FILA_MAX_CARATULA):
        e = norm(ws.cell(r, 1).value)
        if e and e not in etiquetas:
            etiquetas[e] = r
    filas = {}
    for rubro, nombres in RUBROS.items():
        fila = next((etiquetas[n] for n in nombres if n in etiquetas), None)
        if fila is None and rubro not in OPCIONALES:
            alertas.append(f"{unidad}: no se encontró el rubro '{rubro}' ({' / '.join(nombres)})")
        filas[rubro] = fila
    return filas


def leer_unidad(clave, contenido, nombre_archivo, anio, alertas):
    wb = openpyxl.load_workbook(contenido, data_only=True, read_only=False)
    hoja = next((h for h in wb.sheetnames if norm(h) == HOJA), None)
    if not hoja:
        alertas.append(f"{clave}: el archivo {nombre_archivo} no tiene la pestaña '{HOJA}'")
        return None
    ws = wb[hoja]
    encabezados = [norm(c.value) for c in ws[FILA_ENCABEZADOS]]
    cols = columnas_por_mes(encabezados, anio)
    filas = filas_rubros(ws, alertas, clave)

    def valores(col):
        if not col:
            return None
        return {r: (numero(ws.cell(f, col).value) if f else None) for r, f in filas.items()}

    meses = {}
    for m in MESES:
        if m not in cols:
            continue
        real = valores(cols[m]["real"])
        if not real or not real.get("ventas"):
            continue      # mes sin cierre todavía
        ppto = valores(cols[m]["presupuesto"])
        ant = valores(cols[m]["anterior"])
        if ant and not ant.get("ventas"):
            ant = None    # la unidad no operaba ese mes del año anterior
        meses[m] = {
            "real": completar(real),
            "presupuesto": completar(ppto) if ppto and ppto.get("ventas") else None,
            "anio_anterior": completar(ant) if ant else None,
        }
    acu = columnas_acumulado(encabezados, anio)
    acumulado = {"real": completar(valores(acu["real"])) if acu["real"] else None,
                 "anio_anterior": None}
    if acu["anterior"]:
        a = valores(acu["anterior"])
        acumulado["anio_anterior"] = completar(a) if a and a.get("ventas") else None

    return {
        "unidad": UNIDADES[clave]["nombre"],
        "apertura": UNIDADES[clave]["apertura"],
        "archivo_fuente": nombre_archivo,
        "meses": meses,
        "acumulado": acumulado,
    }


def completar(v):
    """Agrega indicadores calculados. Nunca usa el renglón PROMEDIO del Excel."""
    if not v:
        return v
    ventas, com = v.get("ventas"), v.get("comensales")
    v["cheque_promedio"] = round(ventas / com, 2) if ventas and com else None
    if ventas:
        v["margen"] = round((v.get("utilidad") or 0) / ventas, 4)
        v["costo_mas_nomina_pct"] = round(((v.get("costo_venta") or 0) + (v.get("nomina") or 0)) / ventas, 4)
    return v


# ---------------------------------------------------------------- validación

def validar(unidades, alertas):
    for u in unidades.values():
        for m, d in u["meses"].items():
            r = d["real"]
            if r.get("ventas") and not r.get("comensales"):
                alertas.append(f"{u['unidad']} {m}: ventas sin comensales capturados")
            partes = sum((r.get(k) or 0) for k in ("ventas_alimentos", "ventas_bebidas", "otros_ingresos"))
            if r.get("ventas") and abs(partes - r["ventas"]) > 5:
                alertas.append(f"{u['unidad']} {m}: alimentos + bebidas + otros ({partes:,.0f}) "
                               f"no cuadra con total de ingresos ({r['ventas']:,.0f})")


# ---------------------------------------------------------------- principal

def ejecutar(carpeta_local=None):
    alertas = []
    if carpeta_local:
        nombres = [n for n in os.listdir(carpeta_local) if n.lower().endswith(".xlsx")]
        abrir = lambda n: open(os.path.join(carpeta_local, n), "rb")
        origen = f"local:{carpeta_local}"
    else:
        servicio = servicio_drive()
        archivos = listar_drive(servicio, CARPETA_EF_ID)
        por_nombre = {a["name"]: a for a in archivos}
        nombres = list(por_nombre)
        abrir = lambda n: descargar_drive(servicio, por_nombre[n]["id"])
        origen = f"drive:{CARPETA_EF_ID}"

    elegidos, ignorados = elegir_mas_recientes(nombres)
    if not elegidos:
        sys.exit("No se encontró ningún archivo por unidad en la carpeta.")
    for clave in UNIDADES:
        if clave not in elegidos:
            alertas.append(f"{UNIDADES[clave]['nombre']}: no hay archivo en la carpeta")

    anio = max(clasificar(n)[2] for n in elegidos.values())
    unidades = {}
    for clave, nombre in elegidos.items():
        with abrir(nombre) as contenido:
            datos = leer_unidad(clave, contenido, nombre, clasificar(nombre)[2], alertas)
        if datos:
            unidades[clave] = datos

    validar(unidades, alertas)
    ultimo_mes = max((MESES.index(m) for u in unidades.values() for m in u["meses"]), default=None)
    # Periodos en orden cronológico (enero → último cierre) + acumulado.
    # El tablero construye el selector con esta lista y preselecciona el último cierre.
    cerrados = sorted({MESES.index(m) for u in unidades.values() for m in u["meses"]})
    periodos = [{"orden": i + 1, "clave": f"{anio}-{i + 1:02d}", "mes": MESES[i],
                 "etiqueta": f"{MESES[i].capitalize()} {anio}", "tipo": "mes",
                 "predeterminado": i == ultimo_mes} for i in cerrados]
    if cerrados:
        periodos.append({"orden": len(periodos) + 1, "clave": f"{anio}-ACU", "mes": None,
                         "etiqueta": f"Acumulado {MESES[cerrados[0]][:3].lower()}–"
                                     f"{MESES[cerrados[-1]][:3].lower()} {anio}",
                         "tipo": "acumulado", "predeterminado": False})
    for u in unidades.values():   # meses de cada unidad también en orden cronológico
        u["meses"] = {m: u["meses"][m] for m in MESES if m in u["meses"]}
    salida = {
        "generado": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "origen": origen,
        "anio": anio,
        "ultimo_mes_cerrado": MESES[ultimo_mes] if ultimo_mes is not None else None,
        "periodos": periodos,
        "archivos_leidos": {UNIDADES[k]["nombre"]: v for k, v in elegidos.items()},
        "unidades": unidades,
        "alertas": alertas,
    }
    with open(SALIDA, "w", encoding="utf-8") as f:
        json.dump(salida, f, ensure_ascii=False, indent=1)
    print(f"{SALIDA}: {len(unidades)} unidades · último mes cerrado {salida['ultimo_mes_cerrado']} "
          f"· {len(alertas)} alertas")
    for a in alertas:
        print("  ⚠", a)
    return salida


if __name__ == "__main__":
    local = sys.argv[sys.argv.index("--local") + 1] if "--local" in sys.argv else None
    ejecutar(local)
