"""
Pestaña RENTABILIDAD del Tablero Ejecutivo — Grupo Gastronómico Argentilia.

Fuente: estados de resultados de Contraloría, un Excel por restaurante y mes,
cargados en la carpeta de Drive EF_FOLDER_ID con el nombre:

    <UNIDAD> <MES> <AÑO>.xlsx        ej. "ARGENTILIA LEON AGOSTO 2026.xlsx"

Cada archivo trae TODOS los meses del año hasta su cierre (bloques por mes en la
hoja ESTADO DE RESULTADOS), por eso solo se lee el archivo más reciente de cada
unidad. La cifra oficial de cada mes es la columna "TOTAL <MES> <AÑO>".

Archivos que no siguen el patrón (p. ej. los consolidados "AGOSTO 2026.xlsx")
se ignoran.

Acceso a Drive: cuenta de servicio de Google (secreto GDRIVE_SA_KEY con el JSON
de la llave). Para pruebas locales se puede usar EF_LOCAL_DIR con los archivos.
Si no hay acceso, la pestaña muestra un aviso y el resto del tablero se publica
normalmente.
"""
import os
import re
import json
import html as _html
import unicodedata
from io import BytesIO
from datetime import date

import openpyxl

# ── Configuración ──────────────────────────────────────────────────────
EF_FOLDER_ID = os.environ.get("EF_FOLDER_ID", "1jDKGW4dhuy9JKvZ2g8Zuwx17qeE7kBcT")

# Prefijo del nombre de archivo (sin acentos) → nombre de la unidad en el tablero
UNIDADES_EF = {
    "ARGENTILIA LEON":       "Argentilia León",
    "ARGENTILIA QRO":        "Argentilia Querétaro",
    "ARGENTILIA QUERETARO":  "Argentilia Querétaro",
    "FRASCATI":              "Frascati",
    "MIKOH":                 "Mikoh",
    "MANOLO":                "Manolo",
}
NOMBRE_CORTO = {
    "Argentilia León": "A. León", "Argentilia Querétaro": "A. Querétaro",
    "Frascati": "Frascati", "Mikoh": "Mikoh", "Manolo": "Manolo",
}
NOMBRE_TABLA = {"Argentilia León": "León", "Argentilia Querétaro": "Qro",
                "Frascati": "Frascati", "Mikoh": "Mikoh", "Manolo": "Manolo"}

# Fechas de apertura: una unidad es "comparable" en un mes solo si ya operaba
# el mes completo del año anterior. Las unidades no listadas se consideran
# con operación previa a 2025.
APERTURAS = {
    "Mikoh":  date(2025, 9, 25),
    "Manolo": date(2026, 7, 30),
}

# Umbrales del semáforo de estado por unidad
MARGEN_SANO    = 0.18   # margen ≥ 18%            → Sano
MARGEN_CRITICO = 0.05   # margen < 5% o venta cerca/por debajo del equilibrio → Crítico
CUMPL_POR_VALIDAR = 1.50  # cumplimiento > 150% → presupuesto por validar (⚑)
PTS_RUBRO_ALTO = 0.03   # rubro ≥ 3 pts sobre el grupo → contorno rojo

# Responsables y fechas de las palancas (se llenan cuando se definan).
# Claves: 'brecha_venta', 'nomina', 'estructural_<Unidad>', 'gastos', 'bebidas'
RESPONSABLES_PALANCAS = {
    # "nomina": ("Héctor Vázquez", "30/11/2026"),
}

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
         "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
MESES_TIT = [m.capitalize() for m in MESES]
MESES_ABR = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]

# Renglones del estado de resultados (etiqueta normalizada, coincidencia por prefijo)
RENGLONES = {
    "alimentos":     ["INGRESOS ALIMENTOS"],
    "bebidas":       ["INGRESOS BEBIDAS"],
    "ventas":        ["TOTAL DE INGRESOS"],
    "costo":         ["COSTO TOTAL"],
    "nomina":        ["TOTAL NOMINA Y RELATIVOS", "TOTAL NOM"],
    "gastos":        ["TOTAL GASTOS DEPARTAMENTALES"],
    "utilidad":      ["UTILIDAD DESPUES DE OTROS GASTOS"],
    "comensales":    ["COMENSALES"],
    "colaboradores": ["COLABORADORES"],
}
CAMPOS_FIN = ["alimentos", "bebidas", "ventas", "costo", "nomina", "gastos", "utilidad"]


def _norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().upper()


def _num(v):
    try:
        if v is None or v == "":
            return 0.0
        return float(v)
    except (TypeError, ValueError):
        return 0.0


# ── Obtención de archivos ──────────────────────────────────────────────
def interpretar_nombre(nombre):
    """'ARGENTILIA LEON AGOSTO 2026.xlsx' → ('Argentilia León', 8, 2026) o None."""
    base = _norm(re.sub(r"\.xlsx?$", "", nombre, flags=re.I))
    m = re.match(r"^(?P<u>.+?) (?P<mes>" + "|".join(MESES) + r") (?P<anio>\d{4})$", base)
    if not m:
        return None
    unidad = UNIDADES_EF.get(m.group("u"))
    if not unidad:
        return None
    return unidad, MESES.index(m.group("mes")) + 1, int(m.group("anio"))


def _drive_session():
    """Sesión autenticada con la cuenta de servicio (secreto GDRIVE_SA_KEY)."""
    llave = os.environ.get("GDRIVE_SA_KEY", "").strip()
    if not llave:
        return None
    from google.oauth2 import service_account
    from google.auth.transport.requests import AuthorizedSession
    info = json.loads(llave)
    cred = service_account.Credentials.from_service_account_info(
        info, scopes=["https://www.googleapis.com/auth/drive.readonly"])
    return AuthorizedSession(cred)


def descargar_archivo_drive(file_id, sesion=None):
    """Descarga un archivo de Drive con la cuenta de servicio. None si no hay acceso."""
    sesion = sesion or _drive_session()
    if sesion is None:
        return None
    r = sesion.get(f"https://www.googleapis.com/drive/v3/files/{file_id}",
                   params={"alt": "media", "supportsAllDrives": "true"}, timeout=120)
    r.raise_for_status()
    return r.content


def listar_archivos_ef():
    """Lista [(nombre, modificado, cargador)] de la carpeta de Drive o de EF_LOCAL_DIR."""
    local = os.environ.get("EF_LOCAL_DIR")
    if local:
        out = []
        for n in sorted(os.listdir(local)):
            if n.lower().endswith((".xlsx", ".xlsm")):
                ruta = os.path.join(local, n)
                out.append((n, os.path.getmtime(ruta), lambda r=ruta: open(r, "rb").read()))
        return out
    sesion = _drive_session()
    if sesion is None:
        raise RuntimeError("No está configurado el secreto GDRIVE_SA_KEY (cuenta de servicio de Google).")
    archivos, token = [], None
    while True:
        params = {"q": f"'{EF_FOLDER_ID}' in parents and trashed = false",
                  "fields": "nextPageToken, files(id, name, modifiedTime)",
                  "pageSize": 1000, "supportsAllDrives": "true", "includeItemsFromAllDrives": "true"}
        if token:
            params["pageToken"] = token
        r = sesion.get("https://www.googleapis.com/drive/v3/files", params=params, timeout=60)
        r.raise_for_status()
        js = r.json()
        for f in js.get("files", []):
            archivos.append((f["name"], f.get("modifiedTime", ""),
                             lambda fid=f["id"]: descargar_archivo_drive(fid, sesion)))
        token = js.get("nextPageToken")
        if not token:
            break
    return archivos


def seleccionar_ultimos(archivos):
    """Por unidad, el archivo del mes más reciente del año más reciente."""
    elegidos = {}
    for nombre, mod, cargar in archivos:
        info = interpretar_nombre(nombre)
        if not info:
            print(f"   · Se ignora (no es de unidad): {nombre}")
            continue
        unidad, mes, anio = info
        clave = (anio, mes, str(mod))
        if unidad not in elegidos or clave > elegidos[unidad]["clave"]:
            elegidos[unidad] = {"clave": clave, "nombre": nombre, "mes": mes, "anio": anio, "cargar": cargar}
    if not elegidos:
        return {}
    anio_max = max(e["anio"] for e in elegidos.values())
    return {u: e for u, e in elegidos.items() if e["anio"] == anio_max}


# ── Lectura del estado de resultados ───────────────────────────────────
def leer_estado_resultados(contenido, anio, mes_cierre, nombre_archivo=""):
    """Devuelve {'real': {campo: [12]}, 'ppto': {...}, 'ant': {...}} por mes (índice 0 = enero)."""
    wb = openpyxl.load_workbook(BytesIO(contenido), data_only=True, read_only=False)
    hoja = next((ws for ws in wb.worksheets if _norm(ws.title) == "ESTADO DE RESULTADOS"), None)
    if hoja is None:
        raise ValueError(f"{nombre_archivo}: no tiene la hoja 'ESTADO DE RESULTADOS'")

    fila_enc = None
    for r in range(1, min(hoja.max_row, 40) + 1):
        if _norm(hoja.cell(r, 1).value or "") == "DESCRIPCION":
            fila_enc = r
            break
    if fila_enc is None:
        raise ValueError(f"{nombre_archivo}: no se encontró el encabezado DESCRIPCION")

    encab = {c: _norm(hoja.cell(fila_enc, c).value or "") for c in range(1, hoja.max_column + 1)}

    # Columnas por mes: TOTAL <MES> <AÑO> → presupuesto (primer PRESUPUESTO a la derecha) → año anterior (+4)
    cols = {}
    for c in sorted(encab):
        m = re.match(r"^TOTAL (" + "|".join(MESES) + r") (\d{4})$", encab[c])
        if m and int(m.group(2)) == anio:
            idx = MESES.index(m.group(1))
            if idx in cols:
                continue
            c_ppto = next((k for k in range(c + 1, c + 10) if encab.get(k, "").startswith("PRESUPUESTO")), None)
            c_ant = c_ppto + 4 if c_ppto else None
            if c_ant and str(anio - 1) not in encab.get(c_ant, "") and MESES[idx] not in encab.get(c_ant, ""):
                alt = next((k for k in range(c + 1, c + 12) if str(anio - 1) in encab.get(k, "")), None)
                c_ant = alt
            cols[idx] = (c, c_ppto, c_ant)
    faltan = [MESES[i] for i in range(mes_cierre) if i not in cols]
    if faltan:
        raise ValueError(f"{nombre_archivo}: no se encontraron columnas de {', '.join(faltan)}")

    # Renglones (primera aparición)
    filas = {}
    for r in range(fila_enc + 1, hoja.max_row + 1):
        etiqueta = _norm(hoja.cell(r, 1).value or "")
        if not etiqueta:
            continue
        for campo, prefijos in RENGLONES.items():
            if campo not in filas and any(etiqueta == p or etiqueta.startswith(p) for p in prefijos):
                if campo == "nomina" and not etiqueta.startswith("TOTAL NOM"):
                    continue
                filas[campo] = r
    for obligatorio in ("ventas", "costo", "nomina", "gastos", "utilidad"):
        if obligatorio not in filas:
            raise ValueError(f"{nombre_archivo}: no se encontró el renglón '{RENGLONES[obligatorio][0]}'")

    datos = {k: {campo: [0.0] * 12 for campo in RENGLONES} for k in ("real", "ppto", "ant")}
    for idx in range(mes_cierre):
        c_real, c_ppto, c_ant = cols[idx]
        for campo, r in filas.items():
            datos["real"][campo][idx] = _num(hoja.cell(r, c_real).value)
            if c_ppto:
                datos["ppto"][campo][idx] = _num(hoja.cell(r, c_ppto).value)
            if c_ant:
                datos["ant"][campo][idx] = _num(hoja.cell(r, c_ant).value)

    # Validación contra la columna ACUMULADO <AÑO>
    c_acum = next((c for c, t in encab.items() if t == f"ACUMULADO {anio}"), None)
    if c_acum:
        for campo in ("ventas", "utilidad"):
            acum = _num(hoja.cell(filas[campo], c_acum).value)
            suma = sum(datos["real"][campo][:mes_cierre])
            if abs(acum - suma) > max(1000, abs(acum) * 0.005):
                print(f"   ⚠️  {nombre_archivo}: {campo} acumulado {acum:,.0f} ≠ suma de meses {suma:,.0f}")
    return datos


def extraer_financieros():
    """Lee la carpeta y devuelve el modelo base, o {'error': ...} si no se pudo."""
    try:
        archivos = listar_archivos_ef()
    except Exception as e:
        return {"error": f"No fue posible leer la carpeta de estados financieros: {e}"}
    elegidos = seleccionar_ultimos(archivos)
    if not elegidos:
        return {"error": "La carpeta de estados financieros no tiene archivos con el formato 'UNIDAD MES AÑO.xlsx'."}

    unidades, errores = {}, []
    for unidad, e in elegidos.items():
        try:
            datos = leer_estado_resultados(e["cargar"](), e["anio"], e["mes"], e["nombre"])
            datos["mes_cierre"] = e["mes"]
            datos["archivo"] = e["nombre"]
            unidades[unidad] = datos
            print(f"✅ EF {unidad}: {e['nombre']}")
        except Exception as ex:
            errores.append(f"{unidad}: {ex}")
            print(f"⚠️  EF {unidad}: {ex}")
    if not unidades:
        return {"error": "No se pudo leer ningún estado de resultados. " + " | ".join(errores)}
    anio = max(e["anio"] for e in elegidos.values())
    return {"anio": anio, "unidades": unidades, "errores": errores}


# ══════════════════════════════════════════════════════════════════════
#  CÁLCULOS
# ══════════════════════════════════════════════════════════════════════
def _regresion(xs, ys):
    n = len(xs)
    if n < 4:
        return None
    mx, my = sum(xs) / n, sum(ys) / n
    sxx = sum((x - mx) ** 2 for x in xs)
    if sxx <= 0:
        return None
    b = sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / sxx
    return b, my - b * mx


def _div(a, b):
    return a / b if b else None


class Modelo:
    """Métricas por unidad y por grupo para un periodo (lista de índices de mes)."""

    def __init__(self, base):
        self.anio = base["anio"]
        self.u = base["unidades"]
        cierres = [d["mes_cierre"] for d in self.u.values()]
        self.mes_grupo = min(cierres)          # último mes con TODAS las unidades cargadas
        self.mes_max = max(cierres)
        self.pendientes = [u for u, d in self.u.items() if d["mes_cierre"] < self.mes_max]
        self.errores = base.get("errores", [])
        self.orden = [n for n in ["Argentilia León", "Argentilia Querétaro", "Frascati", "Mikoh", "Manolo"] if n in self.u] + \
                     [n for n in self.u if n not in NOMBRE_CORTO]
        # Punto de equilibrio mensual y utilidad marginal por unidad (regresión ene–mes de cierre)
        self.pe, self.pendiente = {}, {}
        for n, d in self.u.items():
            xs, ys = [], []
            for i in range(self.mes_grupo):
                if d["real"]["ventas"][i] > 0:
                    xs.append(d["real"]["ventas"][i]); ys.append(d["real"]["utilidad"][i])
            reg = _regresion(xs, ys)
            if reg and 0 < reg[0] < 1.5:
                b, a = reg
                self.pendiente[n] = b
                self.pe[n] = -a / b if a < 0 else None

    # ── utilidades ──
    def comparable_mes(self, n, i):
        ap = APERTURAS.get(n)
        if ap and not ap < date(self.anio - 1, i + 1, 1):
            return False
        return self.u[n]["ant"]["ventas"][i] > 0

    def comparable(self, n, meses):
        return all(self.comparable_mes(n, i) for i in meses)

    def suma(self, n, tipo, campo, meses):
        return sum(self.u[n][tipo][campo][i] for i in meses)

    def meses_activos(self, n, hasta):
        """Meses completos de operación al cierre del mes índice hasta-1."""
        ap = APERTURAS.get(n)
        if ap:
            return max(1, (self.anio * 12 + hasta - 1) - (ap.year * 12 + ap.month - 1))
        return sum(1 for i in range(hasta) if self.u[n]["real"]["ventas"][i] > 0)

    def metricas(self, n, meses):
        s = lambda t, c: self.suma(n, t, c, meses)
        V, U, C, N, G = (s("real", c) for c in ("ventas", "utilidad", "costo", "nomina", "gastos"))
        V25, U25, C25, N25, G25 = (s("ant", c) for c in ("ventas", "utilidad", "costo", "nomina", "gastos"))
        # presupuesto solo en meses que sí lo tienen capturado
        con_ppto = [i for i in meses if self.u[n]["ppto"]["ventas"][i] > 0]
        Vp = sum(self.u[n]["ppto"]["ventas"][i] for i in con_ppto)
        Up = s("ppto", "utilidad")   # presupuesto de utilidad tal como viene en el archivo (oficial)
        V_eff = sum(self.u[n]["real"]["ventas"][i] for i in con_ppto)
        K, K25 = s("real", "comensales"), s("ant", "comensales")
        m = dict(nombre=n, V=V, U=U, C=C, N=N, G=G, R=V - C - N - G - U,
                 V25=V25, U25=U25, C25=C25, N25=N25, G25=G25, R25=V25 - C25 - N25 - G25 - U25,
                 Vp=Vp, Up=Up, K=K, K25=K25, ali=s("real", "alimentos"), beb=s("real", "bebidas"),
                 activo=V > 0, comparable=self.comparable(n, meses) and V > 0)
        m["margen"] = _div(U, V)
        m["margen25"] = _div(U25, V25)
        m["cumpl"] = _div(V_eff, Vp)
        m["cn"] = _div(C + N, V)
        m["cn25"] = _div(C25 + N25, V25)
        m["cheque"] = _div(V, K)
        m["cheque25"] = _div(V25, K25)
        m["pe"] = self.pe.get(n) if len(meses) == 1 else None
        m["flag_ppto"] = bool(m["cumpl"] and m["cumpl"] > CUMPL_POR_VALIDAR)
        mg = m["margen"] if m["margen"] is not None else 0
        cerca_pe = m["pe"] and V < m["pe"] * 1.05
        m["estado"] = "Crítico" if (mg < MARGEN_CRITICO or cerca_pe) else ("Sano" if mg >= MARGEN_SANO else "Vigilar")
        return m

    def grupo(self, meses, solo_comparables=False):
        ms = [self.metricas(n, meses) for n in self.orden]
        ms = [m for m in ms if m["activo"] and (m["comparable"] or not solo_comparables)]
        g = {k: sum(m[k] for m in ms) for k in ("V", "U", "C", "N", "G", "R", "V25", "U25", "C25", "N25",
                                               "G25", "R25", "Vp", "Up", "K", "K25", "ali", "beb")}
        # presupuesto efectivo: solo meses con presupuesto capturado
        V_eff = sum(m["cumpl"] * m["Vp"] for m in ms if m["cumpl"] is not None)
        g.update(nombre="Grupo", unidades=[m["nombre"] for m in ms],
                 margen=_div(g["U"], g["V"]), margen25=_div(g["U25"], g["V25"]),
                 cumpl=_div(V_eff, g["Vp"]), cn=_div(g["C"] + g["N"], g["V"]),
                 cn25=_div(g["C25"] + g["N25"], g["V25"]), cheque=_div(g["V"], g["K"]),
                 cheque25=_div(g["V25"], g["K25"]))
        return g

    def avisos_datos(self):
        avisos = []
        for n in self.orden:
            d = self.u[n]
            sin = [MESES_TIT[i] for i in range(self.mes_grupo)
                   if d["real"]["ventas"][i] > 0 and d["ppto"]["ventas"][i] <= 0]
            if sin:
                up = sum(d["ppto"]["utilidad"][MESES_TIT.index(x)] for x in sin)
                avisos.append(f"{n}: el archivo no trae presupuesto de venta en {', '.join(sin)} "
                              f"(presupuesto de utilidad {fM(up)}). Ese mes se excluye del cumplimiento de ventas, "
                              f"pero su presupuesto de utilidad sí se suma al acumulado, tal como viene del archivo. "
                              f"Conviene confirmarlo con Contraloría.")
        for e in self.errores:
            avisos.append(f"No se pudo leer: {e}")
        return avisos


# ══════════════════════════════════════════════════════════════════════
#  FORMATO
# ══════════════════════════════════════════════════════════════════════
ROJO, ROJO_OSC, VERDE, AMBAR, AMBAR_OSC = "#C41E28", "#9E1820", "#2E6B4F", "#C99A2E", "#8A6A1C"
GRIS, GRIS_MED, GRIS_CLARO, TINTA, OSCURO = "#656266", "#77727A", "#B5B0AD", "#2C2C2C", "#3F3C40"


def esc(s):
    return _html.escape(str(s))


def fM(x, dec=2):
    if x is None:
        return "n/d"
    return ("−" if x < 0 else "") + f"${abs(x) / 1e6:,.{dec}f}M"


def fP(x, dec=1):
    return "n/d" if x is None else (("−" if x < 0 else "") + f"{abs(x) * 100:,.{dec}f}%")


def fN(x):
    return "n/d" if x is None else f"{x:,.0f}"


def fD(x):
    return "n/d" if x is None else f"${x:,.0f}"


def var(a, b):
    """Variación relativa a vs b."""
    return (a / b - 1) if (a is not None and b) else None


def delta(x, bueno_si_sube=True, pts=False, flecha=True, dec=1):
    """Etiqueta coloreada ▲/▼ para una variación (fracción) o diferencia en puntos."""
    if x is None:
        return f'<span style="color:{GRIS_MED}">n/d</span>'
    neutro = abs(x) < 0.0005
    color = GRIS if neutro else (VERDE if (x > 0) == bueno_si_sube else ROJO)
    signo = "" if neutro else ("+" if x > 0 else "−")
    txt = f"{abs(x) * 100:.{dec}f} pts" if pts else f"{abs(x) * 100:.{dec}f}%"
    if flecha and not neutro:
        return f'<span class="rt-d" style="color:{color}">{"▲" if x > 0 else "▼"} {txt}</span>'
    return f'<span class="rt-d" style="color:{color}">{signo}{txt}</span>'


def lista_es(items):
    items = list(items)
    if len(items) <= 1:
        return "".join(items)
    return ", ".join(items[:-1]) + " y " + items[-1]


# ══════════════════════════════════════════════════════════════════════
#  TEXTOS AUTOMÁTICOS (se regeneran con cada cierre)
# ══════════════════════════════════════════════════════════════════════
def _rubro_mas_alto(m, g):
    """Rubro de costo con mayor exceso sobre el grupo (en % de venta)."""
    rubros = [("nómina", "N"), ("costo de venta", "C"), ("gastos operativos", "G")]
    mejor = None
    for nombre, k in rubros:
        if m["V"] and g["V"]:
            exceso = m[k] / m["V"] - g[k] / g["V"]
            if mejor is None or exceso > mejor[1]:
                mejor = (nombre, exceso, m[k] / m["V"])
    return mejor


def lectura_consejo(mod, meses, ms, g):
    items = []
    # A) Unidades en estado crítico
    for m in ms:
        if m["estado"] != "Crítico":
            continue
        n, V, pe = m["nombre"], m["V"], m["pe"]
        if pe and abs(V / pe - 1) <= 0.05:
            titulo = f"{n} opera en su punto de equilibrio"
        elif pe and V < pe:
            titulo = f"{n} opera por debajo de su punto de equilibrio"
        elif (m["margen"] or 0) < 0:
            titulo = f"{n} opera con pérdida"
        else:
            titulo = f"{n} con margen crítico"
        partes = []
        if pe:
            partes.append(f"Vendió {fM(V)} con un equilibrio estimado de {fM(pe)}.")
        else:
            partes.append(f"Vendió {fM(V)} con margen de {fP(m['margen'])}.")
        frase = ""
        if m["cumpl"] is not None and m["cumpl"] < 0.97:
            frase = f"Está {fP(1 - m['cumpl'], 0)} bajo presupuesto"
        r = _rubro_mas_alto(m, g)
        if r and r[1] > 0.01:
            frase = (frase + " y su " if frase else "Su ") + f"{r[0]} pesa {fP(r[2])} de la venta"
        if frase:
            partes.append(frase + ".")
        items.append((ROJO, titulo, " ".join(partes)))

    # B) Unidades maduras (comparables, no críticas)
    maduras = [m for m in ms if m["comparable"] and m["estado"] != "Crítico"]
    if maduras:
        cumpls = [m["cumpl"] for m in maduras if m["cumpl"] is not None]
        cant = {1: "La unidad madura", 2: "Las dos unidades maduras", 3: "Las tres unidades maduras",
                4: "Las cuatro unidades maduras"}.get(len(maduras), f"Las {len(maduras)} unidades maduras")
        if cumpls:
            lo, hi = min(cumpls), max(cumpls)
            rango = f"{lo * 100:.0f}%" if round(lo * 100) == round(hi * 100) else f"entre {lo * 100:.0f}% y {hi * 100:.0f}%"
            titulo = f"{cant}, {rango} del presupuesto"
        else:
            titulo = cant
        partes = []
        pend = [mod.pendiente[m["nombre"]] for m in maduras if m["nombre"] in mod.pendiente]
        todas_rent = all((m["margen"] or 0) > 0 for m in maduras)
        if cumpls and max(cumpls) < 1:
            base = "Rentables, sin alcanzar su meta de venta" if todas_rent else "Sin alcanzar su meta de venta"
        elif cumpls and min(cumpls) >= 1:
            base = "Rentables y por encima de su meta de venta" if todas_rent else "Por encima de su meta de venta"
        else:
            base = "Cumplimiento mixto contra presupuesto"
        if pend:
            lo, hi = min(pend), max(pend)
            rng = f"${lo:.2f}" if abs(hi - lo) < 0.005 else f"${lo:.2f}–${hi:.2f}"
            base += f"; cada peso adicional deja {rng} de utilidad"
        partes.append(base + ".")
        frases = []
        for m in maduras:
            dk, dc = var(m["K"], m["K25"]), var(m["cheque"], m["cheque25"])
            if dk is None or dc is None or abs(dk) < 0.03:
                continue
            nm = NOMBRE_TABLA.get(m["nombre"], m["nombre"])
            fk, fc = f"{dk * 100:+.1f}%".replace("-", "−"), f"{dc * 100:+.1f}%".replace("-", "−")
            if dk < 0 and dc > 0:
                frases.append(f"{nm} pierde afluencia ({fk} comensales) aunque sube su cheque ({fc})")
            elif dk > 0 and dc < 0:
                frases.append(f"{nm} gana afluencia ({fk}) y baja su cheque ({fc})")
            elif dk < 0:
                frases.append(f"{nm} pierde afluencia ({fk}) y cheque ({fc})")
            else:
                frases.append(f"{nm} gana afluencia ({fk}) y cheque ({fc})")
        if frases:
            partes.append("; ".join(frases) + ".")
        items.append((GRIS_CLARO, titulo, " ".join(partes)))

    # C) Unidades nuevas / referencia de margen
    no_crit = [m for m in ms if m["estado"] != "Crítico" and m["margen"] is not None]
    mejor = max(no_crit, key=lambda m: m["margen"]) if no_crit else None
    nuevas = [m for m in ms if not m["comparable"] and m["estado"] != "Crítico"]
    for m in nuevas:
        n = m["nombre"]
        es_mejor = mejor is not None and mejor["nombre"] == n
        partes = []
        if es_mejor:
            menor_nom = min(ms, key=lambda x: x["N"] / x["V"] if x["V"] else 9)["nombre"] == n
            txt = f"Mejor margen ({fP(m['margen'])})"
            if menor_nom:
                txt += f" y menor peso de nómina ({fP(m['N'] / m['V'])} de la venta)"
            partes.append(txt + ".")
        else:
            partes.append(f"Margen de {fP(m['margen'])}.")
        nmeses = mod.meses_activos(n, max(meses) + 1)
        etapa = "Es su primer mes" if nmeses <= 1 else f"Lleva {nmeses} meses de operación"
        partes.append(etapa + (" y su presupuesto requiere validación." if m["flag_ppto"] else "."))
        titulo = f"{n}: referente, con reserva" if es_mejor else f"{n}: etapa de arranque"
        items.append((AMBAR if (m["flag_ppto"] or es_mejor) else GRIS_CLARO, titulo, " ".join(partes)))
    if not nuevas and mejor is not None:
        n = mejor["nombre"]
        items.append((VERDE, f"{n}: la referencia del grupo",
                      f"Mejor margen ({fP(mejor['margen'])}) con costo + nómina de {fP(mejor['cn'])} de la venta."))
    return items[:4]


def palancas(mod, meses, ms, g):
    """Palancas de utilidad con impacto mensual estimado."""
    out = []
    mes_idx = max(meses)
    n_meses = mod.mes_grupo
    maduras = [m for m in ms if m["comparable"] and m["estado"] != "Crítico" and not m["flag_ppto"]]

    # 1. Brecha de venta
    bajo = [m for m in maduras if m["cumpl"] is not None and m["cumpl"] < 1 and m["Vp"] > 0]
    if bajo:
        imp = 0.0
        for m in bajo:
            pend = mod.pendiente.get(m["nombre"]) or max(0.0, 1 - (m["C"] / m["V"]))
            imp += max(0.0, m["Vp"] - m["V"]) * pend / len(meses)
        out.append(dict(clave="brecha_venta", titulo="Cerrar brecha de venta vs presupuesto",
                        unidades=", ".join(NOMBRE_TABLA.get(m["nombre"], m["nombre"]) for m in bajo),
                        impacto=imp, indicador="Cumplimiento semanal de presupuesto"))

    # 2. Productividad de nómina (comparables, acumulado)
    acum = list(range(mes_idx + 1))
    gc = mod.grupo(acum, solo_comparables=True)
    if gc["V"] and gc["V25"]:
        n26, n25 = gc["N"] / gc["V"], gc["N25"] / gc["V25"]
        if n26 - n25 >= 0.01:
            imp = (n26 - n25) * gc["V"] / len(acum)
            nombres = ", ".join(NOMBRE_TABLA.get(u, u) for u in gc["unidades"])
            out.append(dict(clave="nomina", titulo=f"Productividad de nómina (volver a {fP(n25)})",
                            unidades=nombres, impacto=imp, indicador="Nómina como % de la venta"))

    # 3. Decisión estructural para unidades críticas
    for m in ms:
        if m["estado"] == "Crítico":
            pe = mod.pe.get(m["nombre"])
            out.append(dict(clave=f"estructural_{m['nombre']}", titulo=f"Decisión estructural {m['nombre']}",
                            unidades=m["nombre"], impacto=None, impacto_txt="Evitar pérdida", critica=True,
                            indicador=f"Venta vs equilibrio ({fM(pe)})" if pe else "Margen mensual ≥ 5%"))

    # 4. Gastos operativos por encima del grupo
    if g["V"]:
        g_pct = g["G"] / g["V"]
        altos = [m for m in ms if m["estado"] != "Crítico" and m["V"] and m["G"] / m["V"] - g_pct >= 0.02]
        if altos:
            imp = sum((m["G"] / m["V"] - g_pct) * m["V"] for m in altos) / len(meses)
            out.append(dict(clave="gastos", titulo="Control de gastos operativos",
                            unidades=", ".join(NOMBRE_TABLA.get(m["nombre"], m["nombre"]) for m in altos),
                            impacto=imp, indicador=f"Gastos operativos ≤ {fP(g_pct)} de la venta"))

    # 5. Mezcla de bebidas
    mix = [(m, m["beb"] / m["V"]) for m in ms if m["V"]]
    if len(mix) >= 2:
        ref, pref = max(mix, key=lambda t: t[1])
        rezagadas = [(m, p) for m, p in mix if m is not ref and pref - p >= 0.08]
        if rezagadas:
            ps = [p for _, p in rezagadas]
            out.append(dict(clave="bebidas", titulo=f"Mezcla de bebidas (modelo {ref['nombre']})",
                            unidades="Todas" if len(rezagadas) == len(mix) - 1 else ", ".join(
                                NOMBRE_TABLA.get(m["nombre"], m["nombre"]) for m, _ in rezagadas),
                            impacto=None, impacto_txt="Cheque promedio",
                            indicador=f"% bebidas en venta (hoy {min(ps) * 100:.0f}–{max(ps) * 100:.0f}%)"))
    for i, p in enumerate(out, 1):
        p["num"] = i
        resp = RESPONSABLES_PALANCAS.get(p["clave"])
        p["responsable"], p["fecha"] = (resp if resp else ("[Por definir]", "[Por definir]"))
    return out


def texto_puente(gc, anio):
    dV = var(gc["V"], gc["V25"])
    dU = gc["U"] - gc["U25"]
    efectos = [("Ventas", gc["V"] - gc["V25"], None),
               ("La nómina", -(gc["N"] - gc["N25"]), ("N", "N25")),
               ("El costo de venta", -(gc["C"] - gc["C25"]), ("C", "C25")),
               ("Los gastos operativos", -(gc["G"] - gc["G25"]), ("G", "G25")),
               ("Renta y otros", -(gc["R"] - gc["R25"]), ("R", "R25"))]
    if abs(dV or 0) < 0.02:
        v_txt = f"Con ventas prácticamente iguales ({fP(dV)})"
    else:
        v_txt = f"Con ventas {'al alza' if dV > 0 else 'a la baja'} ({'+' if dV > 0 else ''}{fP(dV)})"
    u_txt = f"la utilidad {'cae' if dU < 0 else 'sube'} {fM(abs(dU))}"
    if abs(dU) < 1:
        return f"{v_txt}, la utilidad se mantiene."
    driver = min(efectos, key=lambda e: e[1]) if dU < 0 else max(efectos, key=lambda e: e[1])
    share = driver[1] / dU if dU else 0
    frase = f"{v_txt}, {u_txt}. {driver[0]} explica {share * 100:.0f}% de la {'caída' if dU < 0 else 'mejora'}"
    if driver[2]:
        a, b = driver[2]
        cre = var(gc[a], gc[b])
        frase += (f": {'+' if cre and cre > 0 else ''}{fP(cre)} en el año, de {fP(_div(gc[b], gc['V25']))} "
                  f"a {fP(_div(gc[a], gc['V']))} de la venta.")
    else:
        frase += "."
    return frase


# ══════════════════════════════════════════════════════════════════════
#  HTML
# ══════════════════════════════════════════════════════════════════════
def _sec(num, titulo, pregunta):
    return (f'<div class="rt-sh"><span class="rt-sn">{num:02d}</span><h2>{esc(titulo)}</h2>'
            f'<span class="rt-sq">{esc(pregunta)}</span></div>')


def _card(titulo, valor, lineas, pie, oscura=False):
    ls = "".join(f'<div class="rt-cl">{l}</div>' for l in lineas if l)
    return (f'<div class="rt-card{" dark" if oscura else ""}"><div class="rt-ct">{esc(titulo)}</div>'
            f'<div class="rt-cv">{valor}</div>{ls}{f"<div class=rt-cf>{pie}</div>" if pie else ""}</div>')


def _vs(d, txt):
    return f'{d}<span class="rt-cm">{txt}</span>'


def _periodo_txt(meses, anio, corto=False):
    if len(meses) == 1:
        i = meses[0]
        return f"{MESES_ABR[i]} {anio}" if corto else f"{MESES_TIT[i]} {anio}"
    return f"{MESES_ABR[meses[0]].lower()}–{MESES_ABR[meses[-1]].lower()} {anio}" if corto else \
        f"{MESES_ABR[meses[0]].lower()}–{MESES_ABR[meses[-1]].lower()}"


def seccion_resultado(mod, meses, g, gc, ms):
    anio, es_mes = mod.anio, len(meses) == 1
    p25 = _periodo_txt(meses, anio - 1, corto=True).replace(f" {anio - 1}", f" {anio - 1}")
    if not es_mes:
        p25 = f"{_periodo_txt(meses, anio)} {anio - 1}"
    sufijo = "del mes" if es_mes else "del periodo"
    no_comp = [m for m in ms if not m["comparable"]]
    aperturas = " · ".join(f"{m['nombre']} abrió {APERTURAS[m['nombre']].strftime('%d/%m/%Y')}"
                           for m in no_comp if m["nombre"] in APERTURAS)
    cards = []
    # Ventas
    pie = ""
    if es_mes and meses[0] > 0:
        prev = [meses[0] - 1]
        v_prev = sum(mod.metricas(u, prev)["V"] for u in gc["unidades"])
        dv = var(gc["V"], v_prev)
        if dv is not None:
            pie = f"{'+' if dv > 0 else ''}{fP(dv)} vs {MESES_TIT[prev[0]].lower()} en unidades comparables"
    elif not es_mes:
        pie = f"Promedio mensual {fM(g['V'] / len(meses))}"
    cards.append(_card(f"Ventas {sufijo}", fM(g["V"]), [
        _vs(delta(var(g["V"], g["Vp"])), f"vs presupuesto {fM(g['Vp'])}"),
        _vs(delta(var(gc["V"], gc["V25"])), f"vs {p25} {fM(gc['V25'])} · comparables") if gc["V25"] else None], pie))
    # Utilidad
    pie = f"Comparables {anio}: {fM(gc['U'])}"
    if es_mes:
        serie = [sum(mod.metricas(u, [i])["U"] for u in gc["unidades"]) for i in range(meses[0] + 1)]
        if len(serie) >= 3:
            if serie[-1] >= max(serie):
                pie += " · mejor mes del año en utilidad"
            elif serie[-1] <= min(serie):
                pie += " · mes más bajo del año en utilidad"
    cards.append(_card(f"Utilidad {sufijo}", fM(g["U"]), [
        _vs(delta(var(g["U"], g["Up"])), f"vs presupuesto {fM(g['Up'])}"),
        _vs(delta(var(gc["U"], gc["U25"])), f"vs {p25} {fM(gc['U25'])} · comparables") if gc["U25"] else None], pie))
    # Margen
    acum = list(range(meses[-1] + 1))
    ga, gac = mod.grupo(acum), mod.grupo(acum, solo_comparables=True)
    l2 = None
    if es_mes and gac["margen25"] is not None:
        l2 = _vs(delta((gac["margen"] or 0) - gac["margen25"], pts=True),
                 f"Acumulado {fP(gac['margen'])} vs {fP(gac['margen25'])} {_periodo_txt(acum, anio)} {anio - 1} · comparables")
    pie = f"Grupo incluye unidades nuevas" if no_comp else ""
    if es_mes:
        pie = f"Comparables {anio}: {fP(gc['margen'])} en el mes · grupo acumulado {fP(ga['margen'])}"
    cards.append(_card("Margen", fP(g["margen"]), [
        _vs(delta((gc["margen"] or 0) - (gc["margen25"] or 0), pts=True),
            f"vs {p25} {fP(gc['margen25'])} · comparables") if gc["margen25"] is not None else None, l2], pie))
    # Comensales
    cards.append(_card(f"Comensales {sufijo}", fN(g["K"]), [
        _vs(delta(var(gc["K"], gc["K25"])), f"vs {p25} {fN(gc['K25'])} · comparables") if gc["K25"] else None],
        f"Comparables {anio}: {fN(gc['K'])}" + (f" · {aperturas}" if aperturas else "")))
    # Cheque
    cards.append(_card("Cheque promedio", fD(g["cheque"]), [
        _vs(delta(var(gc["cheque"], gc["cheque25"])), f"vs {p25} {fD(gc['cheque25'])} · comparables")
        if gc["cheque25"] else None], f"Comparables {anio}: {fD(gc['cheque'])} · ventas totales ÷ comensales"))
    # Tarjeta oscura: utilidad acumulada
    if es_mes:
        brecha = ga["Up"] - ga["U"]
        cards.append(_card(f"Utilidad acumulada {_periodo_txt(acum, anio)}", fM(ga["U"]), [
            _vs(delta(var(ga["U"], ga["Up"])), f"vs presupuesto {fM(ga['Up'])}"),
            _vs(delta(var(gac["U"], gac["U25"])), f"vs {_periodo_txt(acum, anio)} {anio - 1} {fM(gac['U25'])} · comparables")
            if gac["U25"] else None],
            f"Comparables {anio}: {fM(gac['U'])}" + (f" · brecha vs presupuesto {fM(brecha)}" if brecha > 0 else ""),
            oscura=True))
    else:
        brecha = g["Up"] - g["U"]
        cards.append(_card("Brecha de utilidad vs presupuesto", fM(max(brecha, 0)), [
            _vs(delta(var(g["U"], g["Up"])), f"utilidad real vs presupuesto {fM(g['Up'])}")],
            f"Utilidad promedio mensual: {fM(g['U'] / len(meses))}", oscura=True))
    return (f'<section class="rt-s">{_sec(1, "Resultado", "¿Cómo vamos?")}'
            f'<div class="rt-cards">{"".join(cards)}</div></section>')


def _svg_mapa(ms, g):
    pts = [m for m in ms if m["cumpl"] is not None and m["margen"] is not None]
    if not pts:
        return '<div class="rt-empty">Sin datos de presupuesto para el periodo.</div>'
    X0, X1, Y0, Y1 = 60, 560, 20, 280
    xmin = min(0.5, min(0.1 * int(m["cumpl"] * 10) for m in pts))
    xmax = 1.1
    ymin = min(-0.05, min(0.05 * ((m["margen"] // 0.05)) for m in pts))
    ymax = max(0.25, max(0.05 * (-(-m["margen"] // 0.05)) for m in pts))
    sx = lambda v: X0 + (min(v, xmax) - xmin) / (xmax - xmin) * (X1 - X0)
    sy = lambda v: Y1 - (v - ymin) / (ymax - ymin) * (Y1 - Y0)
    gm = g["margen"] or 0
    x100, yg = sx(1.0), sy(gm)
    o = [f'<svg viewBox="0 0 650 330" class="rt-svg" role="img" aria-label="Margen contra cumplimiento de presupuesto por unidad">',
         f'<rect x="{x100:.1f}" y="{Y0}" width="{X1 - x100:.1f}" height="{max(0, yg - Y0):.1f}" fill="#EEF5F1"/>',
         f'<rect x="{X0}" y="{yg:.1f}" width="{x100 - X0:.1f}" height="{max(0, Y1 - yg):.1f}" fill="#FBEFEF"/>',
         f'<line x1="{X0}" y1="{Y1}" x2="{X1}" y2="{Y1}" stroke="#B5B0AD"/><line x1="{X0}" y1="{Y0}" x2="{X0}" y2="{Y1}" stroke="#B5B0AD"/>']
    if ymin < 0 < ymax:
        o.append(f'<line x1="{X0}" y1="{sy(0):.1f}" x2="{X1}" y2="{sy(0):.1f}" stroke="#CCCCCC" stroke-dasharray="2 3"/>'
                 f'<text x="{X1 + 4}" y="{sy(0) + 4:.1f}" font-size="11" fill="{GRIS_MED}">0%</text>')
    o.append(f'<line x1="{x100:.1f}" y1="{Y0}" x2="{x100:.1f}" y2="{Y1}" stroke="{GRIS}" stroke-width="1.2" stroke-dasharray="5 4"/>'
             f'<line x1="{X0}" y1="{yg:.1f}" x2="{X1}" y2="{yg:.1f}" stroke="{GRIS}" stroke-width="1.2" stroke-dasharray="5 4"/>'
             f'<text x="{x100 - 4:.1f}" y="14" font-size="11" fill="{GRIS}" text-anchor="end">100% presupuesto</text>'
             f'<text x="{X1 + 4}" y="{yg + 4:.1f}" font-size="11" fill="{GRIS}">Grupo {fP(gm)}</text>'
             f'<text x="{x100 + 6:.1f}" y="{max(Y0 + 14, yg - 6):.1f}" font-size="12" fill="{VERDE}" font-weight="600">Zona objetivo</text>')
    t = xmin
    while t <= xmax + 1e-9:
        o.append(f'<text x="{sx(t):.1f}" y="300" font-size="11" fill="{GRIS_MED}" text-anchor="middle">{t * 100:.0f}%</text>')
        t += 0.2
    o.append(f'<text x="310" y="322" font-size="12" fill="{GRIS}" text-anchor="middle">Cumplimiento de presupuesto de ventas</text>')
    t = ymax
    while t >= ymin - 1e-9:
        o.append(f'<text x="52" y="{sy(t) + 4:.1f}" font-size="11" fill="{GRIS_MED}" text-anchor="end">{t * 100:.0f}%</text>')
        t -= 0.10
    # texto de la zona roja: abajo, salvo que un punto lo tape (entonces arriba o se omite)
    pts_xy = [(sx(m["cumpl"]), sy(m["margen"])) for m in pts]
    zr = None
    for yz in (Y1 - 8, yg + 17):
        if yz - 12 < yg or X0 + 200 > x100:
            continue
        if not any(X0 + 2 < px < X0 + 215 and yz - 22 < py < yz + 12 for px, py in pts_xy):
            zr = yz
            break
    if zr:
        o.append(f'<text x="{X0 + 10}" y="{zr}" font-size="12" fill="{ROJO_OSC}" font-weight="600">Bajo presupuesto y bajo margen</text>')
    # etiquetas sin encimarse: posiciones candidatas alrededor de cada punto
    zona_y = max(Y0 + 14, yg - 6)
    ocupados = [(x100 + 4, x100 + 96, zona_y - 13, zona_y + 3),          # "Zona objetivo"
                (X0 + 6, X0 + 205, zr - 13, zr + 3) if zr else (0, 0, 0, 0),
                (x100 - 110, x100, 2, 18), (X1 + 2, 650, yg - 10, yg + 6)]  # "100% presupuesto", "Grupo"
    etiquetas = [dict(m=m, cx=sx(m["cumpl"]), cy=sy(m["margen"])) for m in pts]
    for e in etiquetas:
        ocupados.append((e["cx"] - 11, e["cx"] + 11, e["cy"] - 11, e["cy"] + 11))

    def caja(e, izq, ly):
        ancho = 7.0 * max(len(NOMBRE_CORTO.get(e["m"]["nombre"], e["m"]["nombre"])) + 5, 13)
        x0, x1 = (e["cx"] - 15 - ancho, e["cx"] - 15) if izq else (e["cx"] + 15, e["cx"] + 15 + ancho)
        return (x0, x1, ly - 14, ly + 17)

    def libre(r, propio):
        if r[0] < 2 or r[1] > 648 or r[2] < 0 or r[3] > 288:
            return False
        return not any(r[0] < o[1] and o[0] < r[1] and r[2] < o[3] and o[2] < r[3]
                       for o in ocupados if o is not propio)

    for e in sorted(etiquetas, key=lambda e: -e["cx"]):
        propio = next(o for o in ocupados if o == (e["cx"] - 11, e["cx"] + 11, e["cy"] - 11, e["cy"] + 11))
        pref_izq = e["cx"] > 330
        cands = [(pref_izq, 0), (not pref_izq, 0), (pref_izq, 30), (not pref_izq, 30),
                 (pref_izq, -30), (not pref_izq, -30), (pref_izq, 52), (not pref_izq, 52)]
        if e["m"]["cumpl"] > xmax:
            cands = [(True, 26), (True, 44)] + cands
        elegido = cands[0]
        for izq, dy in cands:
            if libre(caja(e, izq, e["cy"] + dy), propio):
                elegido = (izq, dy)
                break
        e["izq"], e["ly"] = elegido[0], e["cy"] + elegido[1]
        ocupados.append(caja(e, e["izq"], e["ly"]))

    for e in etiquetas:
        m, cx, cy = e["m"], e["cx"], e["cy"]
        critico = m["estado"] == "Crítico"
        if m["flag_ppto"]:
            o.append(f'<circle cx="{cx:.1f}" cy="{cy:.1f}" r="9" fill="#FFFFFF" stroke="{AMBAR}" stroke-width="3"/>')
        else:
            o.append(f'<circle cx="{cx:.1f}" cy="{cy:.1f}" r="{10 if critico else 9}" fill="{ROJO if critico else OSCURO}"/>')
        lx = cx - 15 if e["izq"] else cx + 15
        anc = "end" if e["izq"] else "start"
        ly = e["ly"]
        if abs(ly - cy) > 8 and m["cumpl"] <= xmax:
            y1, y2 = (cy + 8, ly - 13) if ly > cy else (cy - 8, ly + 17)
            o.append(f'<line x1="{cx + (-5 if e["izq"] else 5):.1f}" y1="{y1:.1f}" x2="{lx:.1f}" y2="{y2:.1f}" stroke="#B5B0AD"/>')
        nombre = NOMBRE_CORTO.get(m["nombre"], m["nombre"]) + (f" ▸ {m['cumpl'] * 100:.0f}%" if m["cumpl"] > xmax else "")
        col_n, col_s = (ROJO_OSC, ROJO_OSC) if critico else ((TINTA, AMBAR_OSC) if m["flag_ppto"] else (TINTA, GRIS))
        sub = f"{m['cumpl'] * 100:.0f}% · {fP(m['margen'])}" if m["cumpl"] <= xmax else f"{fP(m['margen'])} · ⚑"
        o.append(f'<text x="{lx:.1f}" y="{ly - 2:.1f}" font-size="13" fill="{col_n}" font-weight="600" text-anchor="{anc}">{esc(nombre)}</text>'
                 f'<text x="{lx:.1f}" y="{ly + 13:.1f}" font-size="11" fill="{col_s}" text-anchor="{anc}">{sub}</text>')
    o.append("</svg>")
    return "".join(o)


def _celda(v, sub=None, estilo=""):
    s = f'<div>{v}</div>'
    if sub:
        s += f'<div class="rt-sub2">{sub}</div>'
    return f'<td class="rt-num" style="{estilo}">{s}</td>'


def _tabla_unidades(mod, meses, ms, g, gc):
    anio = mod.anio
    col25 = _periodo_txt(meses, anio - 1, corto=True) if len(meses) == 1 else f"{_periodo_txt(meses, anio)} {anio - 1}"
    orden = sorted([m for m in ms if m["comparable"]], key=lambda m: -(m["margen"] or -9)) + \
        [m for m in ms if not m["comparable"]]
    filas = []
    for k, m in enumerate(orden):
        n = m["nombre"]
        ap = APERTURAS.get(n)
        nota_ap = f"abrió {ap.strftime('%d/%m/%Y')}" if ap else "sin operación comparable"
        tit = esc(n) + (f' <span class="rt-ap">({nota_ap})</span>' if not m["comparable"] else "")
        vp_sub = f'{delta(var(m["V"], m["Vp"]) if m["cumpl"] is None else m["cumpl"] - 1, flecha=False)} <span class="rt-gr">real vs ppto</span>'
        if m["flag_ppto"]:
            vp_sub = vp_sub.replace(f'color:{VERDE}', f'color:{AMBAR_OSC}')
        vp = f'{fM(m["Vp"])}{" ⚑" if m["flag_ppto"] else ""}' if m["Vp"] else '<span class="rt-gr">n/d</span>'
        if m["comparable"]:
            c25 = _celda(fM(m["V25"]), delta(var(m["V"], m["V25"]), flecha=False))
            cn_sub = f'<span class="rt-gr">vs {fP(m["cn25"])} {anio - 1} · </span>{delta((m["cn"] or 0) - (m["cn25"] or 0), bueno_si_sube=False, pts=True, flecha=False)}'
            k_sub = f'{delta(var(m["K"], m["K25"]), flecha=False)} <span class="rt-gr">vs {fN(m["K25"])}</span>'
            ch_sub = f'{delta(var(m["cheque"], m["cheque25"]), flecha=False)} <span class="rt-gr">vs {fD(m["cheque25"])}</span>'
        else:
            c25 = _celda('<span class="rt-gr">n/d</span>', f'<span class="rt-gr">{nota_ap}</span>')
            cn_sub = k_sub = ch_sub = f'<span class="rt-gr">{nota_ap}</span>'
        cn_alto = m["cn"] is not None and g["cn"] and m["cn"] - g["cn"] > 0.05
        u_val = fM(m["U"]) if m["U"] >= 0 else f'<span style="color:{ROJO};font-weight:600">{fM(m["U"])}</span>'
        u_sub = (f'{delta(var(m["U"], m["Up"]) if m["Up"] > 0 else None, flecha=False)} <span class="rt-gr">vs ppto {fM(m["Up"])}</span>'
                 if m["Up"] > 0 else f'<span class="rt-gr">ppto {fM(m["Up"])}</span>')
        cn_val = fP(m["cn"]) if not cn_alto else f'<span style="color:{ROJO};font-weight:600">{fP(m["cn"])}</span>'
        est_col = {"Sano": (VERDE, VERDE), "Vigilar": (AMBAR_OSC, AMBAR), "Crítico": (ROJO_OSC, ROJO)}[m["estado"]]
        filas.append(
            f'<tr class="{"rt-alt" if k % 2 else ""}"><td class="rt-un">{tit}</td>'
            f'<td class="rt-num">{fM(m["V"])}</td>{_celda(vp, vp_sub if m["Vp"] else None)}{c25}'
            f'{_celda(u_val, u_sub)}'
            f'<td class="rt-num" style="font-weight:600;{"color:" + ROJO if (m["margen"] or 0) < 0 else ""}">{fP(m["margen"])}</td>'
            f'{_celda(cn_val, cn_sub)}'
            f'{_celda(fN(m["K"]), k_sub)}{_celda(fD(m["cheque"]), ch_sub)}'
            f'<td class="rt-num">{fM(m["pe"], 1) if m["pe"] else "<span class=rt-gr>n/d</span>"}</td>'
            f'<td><span class="rt-est" style="color:{est_col[0]}"><i style="background:{est_col[1]}"></i>{m["estado"]}</span></td></tr>')
    gr = lambda t: '<span class="rt-gr">' + t + '</span>'
    filas.append(
        '<tr class="rt-tot"><td>Grupo</td>'
        + f'<td class="rt-num">{fM(g["V"])}</td>'
        + _celda(fM(g["Vp"]), delta((g["cumpl"] or 1) - 1, flecha=False) + " " + gr("real vs ppto"))
        + _celda(fM(gc["V25"]), delta(var(gc["V"], gc["V25"]), flecha=False) + " " + gr("comp."))
        + _celda(fM(g["U"]), delta(var(g["U"], g["Up"]), flecha=False) + " " + gr(f"vs ppto {fM(g['Up'])}"))
        + f'<td class="rt-num">{fP(g["margen"])}</td>'
        + _celda(fP(g["cn"]), gr(f"vs {fP(gc['cn25'])} {anio - 1} comp."))
        + _celda(fN(g["K"]), delta(var(gc["K"], gc["K25"]), flecha=False) + " " + gr(f"vs {fN(gc['K25'])} comp."))
        + _celda(fD(g["cheque"]), delta(var(gc["cheque"], gc["cheque25"]), flecha=False) + " " + gr(f"vs {fD(gc['cheque25'])} comp."))
        + '<td class="rt-num rt-gr">—</td><td></td></tr>')
    comp_nombres = lista_es([NOMBRE_TABLA.get(u, u) for u in gc["unidades"]])
    nuevas = [m for m in ms if not m["comparable"]]
    txt_nuevas = ""
    if nuevas:
        txt_nuevas = " (" + "; ".join(
            f"{m['nombre']} abrió el {APERTURAS[m['nombre']].strftime('%d/%m/%Y')}" if m["nombre"] in APERTURAS
            else f"{m['nombre']} sin operación comparable" for m in nuevas) + ")"
    nota = ("* Punto de equilibrio mensual estimado a partir de la relación ventas–utilidad de enero a "
            f"{MESES_TIT[mod.mes_grupo - 1].lower()} (inferencia; requiere al menos 4 meses de operación). "
            if len(meses) == 1 else "* El punto de equilibrio es mensual y se muestra solo en la vista de un mes. ")
    if any(m["flag_ppto"] for m in ms):
        nota += "⚑ Presupuesto por validar con Contraloría (cumplimiento mayor a 150%). "
    nota += (f"Cheque promedio = ventas totales ÷ comensales. “Comp.” = comparativo solo con {comp_nombres}, "
             f"las unidades con operación en el mismo periodo de {anio - 1}{txt_nuevas}. "
             f"Estado: Sano = margen ≥ {MARGEN_SANO * 100:.0f}%; Crítico = margen < {MARGEN_CRITICO * 100:.0f}% "
             f"o venta a menos de 5% de su equilibrio; Vigilar = resto.")
    return (f'<div class="rt-tw"><table class="rt-t"><thead><tr><th style="text-align:left">Unidad</th><th>Ventas</th>'
            f'<th>Presupuesto</th><th>{esc(col25)}</th><th>Utilidad</th><th>Margen</th><th>Costo + Nómina</th><th>Comensales</th>'
            f'<th>Cheque promedio</th><th>Punto de equilibrio*</th><th style="text-align:left">Estado</th></tr></thead>'
            f'<tbody>{"".join(filas)}</tbody></table></div><div class="rt-note">{nota}</div>')


def seccion_mapa(mod, meses, ms, g, gc):
    lect = lectura_consejo(mod, meses, ms, g)
    li = "".join(f'<div class="rt-li"><i style="background:{c}"></i><div><div class="rt-lt">{esc(t)}</div>'
                 f'<div class="rt-lb">{esc(b)}</div></div></div>' for c, t, b in lect)
    return (f'<section class="rt-s">{_sec(2, "Mapa de unidades", "¿Dónde está el problema?")}'
            f'<div class="rt-g2"><div class="rt-box"><div class="rt-ct">Margen vs cumplimiento de presupuesto de ventas</div>'
            f'{_svg_mapa(ms, g)}</div><div class="rt-box"><div class="rt-bh"><div class="rt-ct">Lectura para el consejo</div>'
            f'<span class="rt-pill">Inferencia automática</span></div>{li}</div></div>'
            f'<div class="rt-box rt-flush">{_tabla_unidades(mod, meses, ms, g, gc)}</div></section>')


def seccion_estructura(ms, g, meses):
    comps = [("C", "Costo de venta", "#4A474B", "#FFFFFF"), ("N", "Nómina", "#7A767A", "#FFFFFF"),
             ("G", "Gastos operativos", "#B5B0AD", TINTA), ("R", "Renta y otros", "#DAD6D3", TINTA)]
    leyenda = "".join(f'<span><i style="background:{c}"></i>{esc(t)}</span>' for _, t, c, _ in comps) + \
        f'<span><i style="background:{ROJO}"></i>Utilidad</span>'

    def barra(m, es_grupo=False):
        V = m["V"]
        if not V:
            return ""
        segs = []
        for k, _, col, txt in comps:
            p = m[k] / V
            alto = (not es_grupo) and g["V"] and (p - g[k] / g["V"]) >= PTS_RUBRO_ALTO
            ancho = max(p, 0) * 100
            etiqueta = f"{p * 100:.1f}" if ancho >= 5.5 else ""
            segs.append(f'<div style="width:{ancho:.2f}%;background:{col};color:{txt};'
                        f'{"box-shadow:inset 0 0 0 3px #ED2E38;" if alto else ""}">{etiqueta}</div>')
        u = m["U"] / V
        if u > 0:
            segs.append(f'<div style="width:{u * 100:.2f}%;background:{ROJO};color:#FFF">{u * 100:.1f}</div>')
        neg = f'<span class="rt-neg">{u * 100:.1f}</span>'.replace("-", "−") if u <= 0 else ""
        nombre = "Grupo" if es_grupo else NOMBRE_CORTO.get(m["nombre"], m["nombre"])
        crit = (not es_grupo) and m["estado"] == "Crítico"
        return (f'<div class="rt-br{" rt-brg" if es_grupo else ""}"><div class="rt-bn" style="{"color:" + ROJO_OSC if crit else ""}">'
                f'{esc(nombre)}</div><div class="rt-bt"><div class="rt-bb">{"".join(segs)}</div>{neg}</div></div>')

    orden = sorted([m for m in ms if m["V"]], key=lambda m: -(m["margen"] or -9))
    ref = lista_es([NOMBRE_TABLA.get(m["nombre"], m["nombre"]) for m in orden[:2]])
    per = "el mes" if len(meses) == 1 else "el periodo"
    return (f'<section class="rt-s">{_sec(3, "Estructura de costo", "¿Por qué?")}<div class="rt-box rt-eq">'
            f'<div class="rt-ct">Cómo se reparte cada $100 de venta · {per}</div><div class="rt-leg">{leyenda}</div>'
            f'<div class="rt-bars">{"".join(barra(m) for m in orden)}{barra(g, True)}</div>'
            f'<div class="rt-note" style="margin-top:auto">Contorno rojo = rubro {PTS_RUBRO_ALTO * 100:.0f} pts o más por encima del grupo. '
            f'Referencia interna: {esc(ref)}. “Renta y otros” incluye renta, mantenimiento de plaza y gastos varios.</div></div></section>')


def seccion_puente(mod, meses):
    acum = list(range(meses[-1] + 1))
    gc = mod.grupo(acum, solo_comparables=True)
    anio = mod.anio
    titulo_per = _periodo_txt(acum, anio)
    if not gc["unidades"] or not gc["V25"]:
        cuerpo = '<div class="rt-empty">No hay unidades con operación comparable en el año anterior.</div>'
        return f'<section class="rt-s">{_sec(4, "Puente de utilidad", f"¿Qué cambió vs {anio - 1}?")}<div class="rt-box rt-eq">{cuerpo}</div></section>'
    pasos = [("Ventas", gc["V"] - gc["V25"]), ("Costo de|venta", -(gc["C"] - gc["C25"])),
             ("Nómina", -(gc["N"] - gc["N25"])), ("Gastos|operativos", -(gc["G"] - gc["G25"])),
             ("Renta y|otros", -(gc["R"] - gc["R25"]))]
    vals = [gc["U25"]]
    for _, d in pasos:
        vals.append(vals[-1] + d)
    top, bot = max(vals + [0]) / 1e6, min(vals + [0]) / 1e6
    rango = (top - bot) or 1
    Y0, Y1 = 40, 270
    sy = lambda v: Y1 - (v / 1e6 - bot) / rango * (Y1 - Y0)
    peor = min(pasos, key=lambda p: p[1])[0]
    o = [f'<svg viewBox="0 0 620 330" class="rt-svg" role="img" aria-label="Cascada de utilidad {anio - 1} a {anio}">',
         f'<line x1="20" y1="{sy(0):.1f}" x2="610" y2="{sy(0):.1f}" stroke="#B5B0AD"/>']

    def rect(x, v0, v1, color):
        y, h = min(sy(v0), sy(v1)), abs(sy(v0) - sy(v1))
        return f'<rect x="{x}" y="{y:.1f}" width="62" height="{max(h, 1.5):.1f}" fill="{color}"/>'

    def etiqueta(x, lineas, color, bold=False):
        return "".join(f'<text x="{x + 31}" y="{290 + 15 * i}" font-size="12" fill="{color}" text-anchor="middle"'
                       f'{" font-weight=" + chr(34) + "700" + chr(34) if bold else ""}>{esc(t)}</text>' for i, t in enumerate(lineas))

    x = 30
    o.append(rect(x, 0, gc["U25"], GRIS))
    o.append(f'<text x="{x + 31}" y="{sy(max(gc["U25"], 0)) - 9:.1f}" font-size="15" font-weight="700" fill="{TINTA}" text-anchor="middle">{gc["U25"] / 1e6:.2f}</text>')
    o.append(etiqueta(x, [str(anio - 1)], TINTA, True))
    acumv = gc["U25"]
    for nombre, d in pasos:
        o.append(f'<line x1="{x + 62}" y1="{sy(acumv):.1f}" x2="{x + 84}" y2="{sy(acumv):.1f}" stroke="#B5B0AD" stroke-dasharray="3 3"/>')
        x += 84
        grande = nombre == peor and d < 0
        color = (ROJO if grande else "#E88A8F") if d < 0 else "#6FA88A"
        o.append(rect(x, acumv, acumv + d, color))
        ytxt = (sy(min(acumv, acumv + d)) + 18) if d < 0 else (sy(max(acumv, acumv + d)) - 8)
        o.append(f'<text x="{x + 31}" y="{ytxt:.1f}" font-size="{15 if grande else 14}" font-weight="{700 if grande else 600}" '
                 f'fill="{ROJO_OSC if d < 0 else VERDE}" text-anchor="middle">{"+" if d >= 0 else "−"}{abs(d) / 1e6:.2f}</text>')
        o.append(etiqueta(x, nombre.split("|"), ROJO_OSC if grande else GRIS, grande))
        acumv += d
    o.append(f'<line x1="{x + 62}" y1="{sy(acumv):.1f}" x2="{x + 84}" y2="{sy(acumv):.1f}" stroke="#B5B0AD" stroke-dasharray="3 3"/>')
    x += 84
    o.append(rect(x, 0, gc["U"], OSCURO))
    o.append(f'<text x="{x + 31}" y="{sy(max(gc["U"], 0)) - 9:.1f}" font-size="15" font-weight="700" fill="{TINTA}" text-anchor="middle">{gc["U"] / 1e6:.2f}</text>')
    o.append(etiqueta(x, [str(anio)], TINTA, True))
    o.append("</svg>")
    du = var(gc["U"], gc["U25"])
    fondo, col = ("#FDF0F0", ROJO) if (du or 0) < 0 else ("#EEF5F1", VERDE)
    unidades = lista_es([NOMBRE_TABLA.get(u, u) for u in gc["unidades"]])
    return (f'<section class="rt-s">{_sec(4, "Puente de utilidad", f"¿Qué cambió vs {anio - 1}?")}<div class="rt-box rt-eq">'
            f'<div class="rt-bh"><div class="rt-ct">Utilidad acumulada {titulo_per} · {esc(unidades)} (millones)</div>'
            f'<span class="rt-pill">Dato confirmado</span></div>{"".join(o)}'
            f'<div class="rt-call" style="background:{fondo}"><div class="rt-callv" style="color:{col}">'
            f'{"+" if (du or 0) > 0 else ""}{fP(du)}</div><div>{esc(texto_puente(gc, anio))}</div></div></div></section>')


def seccion_palancas(mod, meses, ms, g):
    pal = palancas(mod, meses, ms, g)
    if not pal:
        return ""
    total = sum(p["impacto"] for p in pal if p.get("impacto"))
    criticas = [p["unidades"] for p in pal if p.get("critica")]
    acum = list(range(meses[-1] + 1))
    ga = mod.grupo(acum)
    brecha = ga["Up"] - ga["U"]
    filas = []
    for k, p in enumerate(pal):
        imp = (f'<td class="rt-num" style="font-weight:700;color:{VERDE}">+{fM(p["impacto"])}</td>' if p.get("impacto")
               else f'<td class="rt-num" style="font-weight:600">{esc(p.get("impacto_txt", "—"))}</td>')
        pend = lambda t: f'<span class="rt-gr">{esc(t)}</span>' if t.startswith("[") else esc(t)
        filas.append(f'<tr class="{"rt-alt" if k % 2 else ""}"><td style="font-weight:600;{"color:" + ROJO_OSC if p.get("critica") else ""}">'
                     f'{p["num"]}. {esc(p["titulo"])}</td><td class="rt-gr">{esc(p["unidades"])}</td>{imp}'
                     f'<td>{esc(p["indicador"])}</td><td>{pend(p["responsable"])}</td><td>{pend(p["fecha"])}</td></tr>')
    meses_rec = f"Equivale a recuperar en ~{brecha / total:.0f} meses la brecha acumulada contra presupuesto ({fM(brecha)})." \
        if total > 0 and brecha > 0 else ""
    sin = f"por mes, sin contar {lista_es([c for c in criticas])}" if criticas else "por mes"
    return (f'<section class="rt-s">{_sec(5, "Palancas y plan", "¿Qué vamos a hacer y cuánto vale?")}<div class="rt-g3">'
            f'<div class="rt-card dark rt-pot"><div class="rt-ct">Potencial de utilidad identificado</div>'
            f'<div class="rt-potv">{fM(total, 1)}</div><div class="rt-pots">{esc(sin)}</div><hr>'
            f'<div class="rt-cf">{esc(meses_rec)}</div><span class="rt-pill dark">Estimación por validar</span></div>'
            f'<div class="rt-box rt-flush"><div class="rt-tw"><table class="rt-t"><thead><tr><th style="text-align:left">Palanca</th>'
            f'<th style="text-align:left">Unidades</th><th>Impacto / mes</th><th style="text-align:left">Indicador de seguimiento</th>'
            f'<th style="text-align:left">Responsable</th><th style="text-align:left">Fecha</th></tr></thead><tbody>{"".join(filas)}</tbody></table></div>'
            f'<div class="rt-note">Impactos estimados con la utilidad marginal de cada unidad y la estructura de costo del periodo; '
            f'se recalculan con cada cierre. Responsables y fechas se definen con la dirección.</div></div></div></section>')


def vista(mod, meses):
    ms = [mod.metricas(n, meses) for n in mod.orden]
    ms = [m for m in ms if m["activo"]]
    g, gc = mod.grupo(meses), mod.grupo(meses, solo_comparables=True)
    return (seccion_resultado(mod, meses, g, gc, ms) + seccion_mapa(mod, meses, ms, g, gc) +
            f'<div class="rt-g2 rt-g2b">{seccion_estructura(ms, g, meses)}{seccion_puente(mod, meses)}</div>' +
            seccion_palancas(mod, meses, ms, g))


CSS = """
#rentabilidad .rt{font-family:'IBM Plex Sans','Segoe UI',Arial,sans-serif;color:#2C2C2C;background:#F4F2F0;border-radius:12px;
  padding:28px 32px 36px;font-variant-numeric:tabular-nums;display:flex;flex-direction:column;gap:28px}
#rentabilidad .rt-top{height:4px;background:#ED2E38;border-radius:2px}
#rentabilidad .rt-hd{display:flex;justify-content:space-between;align-items:flex-end;gap:24px;flex-wrap:wrap;margin-top:-10px}
#rentabilidad .rt-kick{font-size:12px;letter-spacing:.14em;text-transform:uppercase;color:#77727A;font-weight:600}
#rentabilidad .rt-h1{margin:4px 0 2px;font-family:'IBM Plex Sans Condensed','IBM Plex Sans',sans-serif;font-size:40px;font-weight:700;color:#3F3C40;letter-spacing:-.01em}
#rentabilidad .rt-hs{font-size:15px;color:#656266}
#rentabilidad .rt-per{display:flex;gap:10px;align-items:center}
#rentabilidad .rt-per label{font-size:13px;color:#77727A}
#rentabilidad .rt-per select{height:42px;padding:0 14px;font:inherit;font-size:14px;border:1px solid #D6D2CF;border-radius:8px;background:#fff;color:#2C2C2C}
#rentabilidad .rt-banner{background:#FEF3E2;border:1px solid #C99A2E;color:#8A6A1C;border-radius:8px;padding:10px 16px;font-size:14px;font-weight:600}
#rentabilidad .rt-error{background:#FDECEA;border:1px solid #ED2E38;color:#9E1820;border-radius:8px;padding:16px 20px;font-size:14px;line-height:1.5}
#rentabilidad .rt-s{display:flex;flex-direction:column;gap:14px;min-width:0}
#rentabilidad .rt-sh{display:flex;align-items:baseline;gap:14px;border-bottom:1px solid #D6D2CF;padding-bottom:10px}
#rentabilidad .rt-sh h2{margin:0;font-size:20px;font-weight:600;color:#3F3C40}
#rentabilidad .rt-sn{font-size:13px;font-weight:700;color:#ED2E38}
#rentabilidad .rt-sq{font-size:15px;color:#77727A}
#rentabilidad .rt-cards{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:16px}
#rentabilidad .rt-card{background:#fff;border-radius:12px;padding:20px 22px;display:flex;flex-direction:column;gap:7px;border:1px solid #E7E4E1}
#rentabilidad .rt-card.dark{background:#3F3C40;border-color:#3F3C40}
#rentabilidad .rt-ct{font-size:14px;color:#656266;font-weight:500}
#rentabilidad .rt-card.dark .rt-ct{color:#E7E4E1}
#rentabilidad .rt-cv{font-size:36px;font-weight:600;color:#2C2C2C;letter-spacing:-.01em}
#rentabilidad .rt-card.dark .rt-cv{color:#fff}
#rentabilidad .rt-cl{display:flex;gap:8px;align-items:baseline;font-size:14px;flex-wrap:wrap}
#rentabilidad .rt-cm{color:#656266}
#rentabilidad .rt-card.dark .rt-cm{color:#E7E4E1}
#rentabilidad .rt-card.dark .rt-d[style*="C41E28"]{color:#FF8A8F!important}
#rentabilidad .rt-card.dark .rt-d[style*="2E6B4F"]{color:#8FD1AE!important}
#rentabilidad .rt-cf{font-size:13px;color:#77727A;line-height:1.45}
#rentabilidad .rt-card.dark .rt-cf{color:#CFCAC7}
#rentabilidad .rt-d{font-weight:600;white-space:nowrap}
#rentabilidad .rt-g2{display:grid;grid-template-columns:minmax(0,1.05fr) minmax(0,1fr);gap:16px}
#rentabilidad .rt-g3{display:grid;grid-template-columns:300px minmax(0,1fr);gap:16px}
#rentabilidad .rt-box{background:#fff;border-radius:12px;padding:20px 24px;border:1px solid #E7E4E1;display:flex;flex-direction:column;gap:14px;min-width:0}
#rentabilidad .rt-box.rt-flush{padding:0;gap:0;overflow:hidden}
#rentabilidad .rt-eq{min-height:470px}
#rentabilidad .rt-bh{display:flex;justify-content:space-between;align-items:center;gap:10px}
#rentabilidad .rt-pill{font-size:12px;padding:4px 10px;border-radius:999px;background:#F1EFED;color:#656266;font-weight:600;white-space:nowrap}
#rentabilidad .rt-pill.dark{background:#56525A;color:#fff;align-self:flex-start}
#rentabilidad .rt-svg{width:100%;height:auto;display:block;font-family:inherit}
#rentabilidad .rt-li{display:flex;gap:14px}
#rentabilidad .rt-li>i{width:6px;border-radius:3px;flex-shrink:0}
#rentabilidad .rt-lt{font-size:16px;font-weight:600;color:#2C2C2C;margin-bottom:3px}
#rentabilidad .rt-lb{font-size:14.5px;color:#656266;line-height:1.5}
#rentabilidad .rt-tw{overflow-x:auto}
#rentabilidad table.rt-t{width:100%;border-collapse:collapse;font-size:13.5px}
#rentabilidad .rt-t th{background:#ED2E38;color:#fff;font-weight:600;padding:12px 9px;text-align:right;font-size:13px;letter-spacing:0;white-space:nowrap}
#rentabilidad .rt-t td{padding:11px 9px;border-bottom:none;vertical-align:top}
#rentabilidad .rt-t tbody tr{background:#FDF0F0}
#rentabilidad .rt-t tbody tr.rt-alt{background:#F5F5F5}
#rentabilidad .rt-t tbody tr.rt-tot{background:#fff;border-top:2px solid #3F3C40;font-weight:700}
#rentabilidad .rt-num{text-align:right;white-space:nowrap}
#rentabilidad .rt-t td:nth-last-child(-n+2){white-space:nowrap}
#rentabilidad .rt-un{font-weight:600;color:#3F3C40}
#rentabilidad .rt-ap{font-weight:400;font-size:12.5px;color:#77727A}
#rentabilidad .rt-sub2{font-size:11.5px;margin-top:2px;font-weight:400}
#rentabilidad .rt-gr{color:#77727A;font-weight:400}
#rentabilidad .rt-est{display:inline-flex;align-items:center;gap:8px;font-size:14px;font-weight:600}
#rentabilidad .rt-est i{width:10px;height:10px;border-radius:50%}
#rentabilidad .rt-note{padding:12px 20px;font-size:12.5px;color:#77727A;border-top:1px solid #E7E4E1;line-height:1.5}
#rentabilidad .rt-box:not(.rt-flush) .rt-note{padding:0;border:0}
#rentabilidad .rt-leg{display:flex;gap:16px;flex-wrap:wrap;font-size:13px;color:#656266}
#rentabilidad .rt-leg span{display:inline-flex;align-items:center;gap:6px}
#rentabilidad .rt-leg i{width:12px;height:12px;border-radius:2px}
#rentabilidad .rt-bars{display:flex;flex-direction:column;gap:13px}
#rentabilidad .rt-br{display:flex;align-items:center;gap:12px}
#rentabilidad .rt-brg{padding-top:8px;border-top:1px dashed #CCCCCC}
#rentabilidad .rt-bn{width:100px;flex-shrink:0;font-size:14px;font-weight:600;color:#3F3C40}
#rentabilidad .rt-brg .rt-bn{font-weight:700;color:#2C2C2C}
#rentabilidad .rt-bt{flex:1;display:flex;align-items:center;gap:8px;min-width:0}
#rentabilidad .rt-bb{display:flex;height:34px;font-size:12px;font-weight:600;width:100%}
#rentabilidad .rt-bb>div{display:flex;align-items:center;justify-content:center;overflow:hidden;white-space:nowrap}
#rentabilidad .rt-neg{font-size:12px;font-weight:700;color:#9E1820;white-space:nowrap}
#rentabilidad .rt-call{display:flex;gap:14px;align-items:flex-start;padding:14px 16px;border-radius:8px;font-size:14px;color:#3F3C40;line-height:1.5}
#rentabilidad .rt-callv{font-size:26px;font-weight:700;line-height:1;white-space:nowrap}
#rentabilidad .rt-pot{justify-content:center;gap:12px;padding:26px}
#rentabilidad .rt-potv{font-size:50px;font-weight:700;color:#fff;letter-spacing:-.02em;line-height:1}
#rentabilidad .rt-pots{font-size:16px;color:#E7E4E1}
#rentabilidad .rt-pot hr{border:0;height:1px;background:#656266;margin:4px 0}
#rentabilidad .rt-empty{color:#77727A;font-size:14px;padding:24px 0}
#rentabilidad .rt-calidad{font-size:12.5px;color:#77727A;line-height:1.5}
#rentabilidad .rt-calidad li{margin-left:18px}
@media (max-width:1150px){
  #rentabilidad .rt-cards{grid-template-columns:repeat(2,minmax(0,1fr))}
  #rentabilidad .rt-g2,#rentabilidad .rt-g3{grid-template-columns:minmax(0,1fr)}
  #rentabilidad .rt-eq{min-height:0}
}
@media (max-width:700px){
  #rentabilidad .rt{padding:18px 16px}
  #rentabilidad .rt-cards{grid-template-columns:minmax(0,1fr)}
  #rentabilidad .rt-h1{font-size:32px}
}
@media print{
  #rentabilidad .rt{background:#fff;padding:0}
  #rentabilidad .rt-per{display:none}
  #rentabilidad .rt-s,#rentabilidad .rt-card,#rentabilidad .rt-box{break-inside:avoid}
}
"""

HEAD = ('<link rel="preconnect" href="https://fonts.googleapis.com">'
        '<link href="https://fonts.googleapis.com/css2?family=IBM+Plex+Sans:wght@400;500;600;700&amp;'
        'family=IBM+Plex+Sans+Condensed:wght@600;700&amp;display=swap" rel="stylesheet">')

JS = """
function rtVer(v){document.querySelectorAll('#rentabilidad [data-rt-vista]').forEach(function(d){
  d.style.display = d.getAttribute('data-rt-vista')===v ? '' : 'none';});
  var s=document.getElementById('rt-sub'); var o=document.querySelector('#rt-periodo option[value="'+v+'"]');
  if(s&&o) s.textContent=o.getAttribute('data-sub');}
"""


def generar_seccion(base, fecha_actualizacion=""):
    """Devuelve el HTML completo de la sección (div#rentabilidad)."""
    cab = ('<div class="rt-top"></div><div class="rt-hd"><div><div class="rt-kick">Estados financieros · Contraloría</div>'
           '<h1 class="rt-h1">Rentabilidad</h1><div class="rt-hs" id="rt-sub">{sub}</div></div>{sel}</div>')
    if not base or base.get("error"):
        msg = esc(base.get("error") if base else "Sin información financiera.")
        return (f'<div id="rentabilidad" class="section"><div class="rt">{cab.format(sub="Sin información disponible", sel="")}'
                f'<div class="rt-error"><strong>La información financiera no está disponible en esta actualización.</strong><br>{msg}<br>'
                f'El resto del tablero se actualizó normalmente.</div></div></div>')
    mod = Modelo(base)
    M = mod.mes_grupo
    vistas = [[i] for i in range(M) if any(mod.u[n]["real"]["ventas"][i] > 0 for n in mod.u)]
    opciones, bloques = [], []
    for meses in reversed(vistas):
        clave = f"m{meses[0] + 1}"
        sub = f"{MESES_TIT[meses[0]]} {mod.anio} · mes cerrado · fuente: estados de resultados de Contraloría"
        opciones.append((clave, f"{MESES_TIT[meses[0]]} {mod.anio}", sub))
        bloques.append((clave, vista(mod, meses)))
    if M > 1:
        acum = list(range(M))
        clave = "acum"
        sub = f"Acumulado {_periodo_txt(acum, mod.anio)} {mod.anio} · meses cerrados · fuente: estados de resultados de Contraloría"
        opciones.insert(1, (clave, f"Acumulado {_periodo_txt(acum, mod.anio)}", sub))
        bloques.insert(1, (clave, vista(mod, acum)))
    defecto = opciones[0][0]
    sel = ('<div class="rt-per"><label for="rt-periodo">Periodo</label><select id="rt-periodo" onchange="rtVer(this.value)">' +
           "".join(f'<option value="{k}" data-sub="{esc(s)}"{" selected" if k == defecto else ""}>{esc(t)}</option>'
                   for k, t, s in opciones) + "</select></div>")
    banner = ""
    if mod.pendientes:
        mes_n = MESES_TIT[mod.mes_max - 1]
        cargadas = len(mod.u) - len(mod.pendientes)
        banner = (f'<div class="rt-banner">{mes_n}: {cargadas} de {len(mod.u)} unidades cargadas '
                  f'(faltan {esc(lista_es(mod.pendientes))}). El tablero muestra {MESES_TIT[M - 1]}, '
                  f'último mes completo del grupo.</div>')
    avisos = mod.avisos_datos()
    calidad = ""
    if avisos:
        calidad = ('<div class="rt-calidad"><strong>Notas de calidad de datos</strong><ul>' +
                   "".join(f"<li>{esc(a)}</li>" for a in avisos) + "</ul></div>")
    archivos = " · ".join(f"{NOMBRE_TABLA.get(n, n)}: {MESES_ABR[mod.u[n]['mes_cierre'] - 1]}" for n in mod.orden)
    pie = (f'<div class="rt-calidad">Archivos leídos (último cierre por unidad): {esc(archivos)}.'
           f'{" Revisión: " + esc(fecha_actualizacion) + "." if fecha_actualizacion else ""} '
           f'Los textos de lectura y las palancas se generan automáticamente con cada cierre y deben validarse con la dirección.</div>')
    cuerpo = "".join(f'<div data-rt-vista="{k}"{"" if k == defecto else " style=" + chr(34) + "display:none" + chr(34)}>'
                     f'<div class="rt" style="padding:0;background:none;gap:28px">{h}</div></div>' for k, h in bloques)
    return (f'<div id="rentabilidad" class="section"><div class="rt">{cab.format(sub=esc(opciones[0][2]), sel=sel)}'
            f'{banner}{cuerpo}{calidad}{pie}</div></div>')
